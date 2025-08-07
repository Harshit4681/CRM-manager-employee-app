import asyncio
import hashlib
import httpx
import logging
import uuid
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Dict, Any
import tempfile
import os

# Document processing imports
import PyPDF2
from docx import Document
import fitz  # PyMuPDF for better PDF handling
from email import message_from_string
from email.header import decode_header
import magic

from app.models.schemas import DocumentInfo, DocumentChunk, DocumentType
from app.core.config import settings
from app.utils.text_processing import TextProcessor

logger = logging.getLogger(__name__)

class DocumentProcessor:
    """Service for processing documents from blob URLs"""
    
    def __init__(self):
        self.text_processor = TextProcessor()
        self.supported_types = {
            'application/pdf': DocumentType.PDF,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document': DocumentType.DOCX,
            'text/plain': DocumentType.TXT,
            'message/rfc822': DocumentType.EMAIL
        }
        
    async def process_document(self, blob_url: str) -> DocumentInfo:
        """
        Process a document from a blob URL
        
        Args:
            blob_url: URL to the document blob
            
        Returns:
            DocumentInfo with processed content and chunks
        """
        try:
            logger.info(f"Starting processing of document: {blob_url}")
            
            # Download document
            file_data, content_type, filename = await self._download_document(blob_url)
            
            # Detect document type
            doc_type = self._detect_document_type(content_type, filename, file_data)
            
            # Generate document ID
            document_id = self._generate_document_id(blob_url, file_data)
            
            # Process based on type
            if doc_type == DocumentType.PDF:
                content, metadata = await self._process_pdf(file_data)
            elif doc_type == DocumentType.DOCX:
                content, metadata = await self._process_docx(file_data)
            elif doc_type == DocumentType.EMAIL:
                content, metadata = await self._process_email(file_data)
            elif doc_type == DocumentType.TXT:
                content, metadata = await self._process_text(file_data)
            else:
                raise ValueError(f"Unsupported document type: {doc_type}")
            
            # Create chunks
            chunks = await self._create_chunks(content, document_id, metadata)
            
            # Build document info
            doc_info = DocumentInfo(
                document_id=document_id,
                filename=filename,
                document_type=doc_type,
                blob_url=blob_url,
                file_size=len(file_data),
                processed_at=datetime.utcnow(),
                chunks=chunks,
                metadata=metadata
            )
            
            logger.info(f"Successfully processed document: {filename} ({len(chunks)} chunks)")
            return doc_info
            
        except Exception as e:
            logger.error(f"Error processing document {blob_url}: {e}")
            raise
    
    async def _download_document(self, blob_url: str) -> tuple[bytes, str, str]:
        """Download document from blob URL"""
        try:
            async with httpx.AsyncClient(timeout=30.0) as client:
                response = await client.get(blob_url)
                response.raise_for_status()
                
                # Get content type
                content_type = response.headers.get('content-type', 'application/octet-stream')
                
                # Extract filename from URL or headers
                filename = self._extract_filename(blob_url, response.headers)
                
                return response.content, content_type, filename
                
        except httpx.RequestError as e:
            raise ValueError(f"Failed to download document: {e}")
    
    def _extract_filename(self, url: str, headers: Dict[str, str]) -> str:
        """Extract filename from URL or headers"""
        # Try content-disposition header first
        content_disposition = headers.get('content-disposition', '')
        if 'filename=' in content_disposition:
            filename = content_disposition.split('filename=')[1].strip('"')
            return filename
        
        # Fall back to URL path
        path = Path(url.split('?')[0])  # Remove query params
        if path.suffix:
            return path.name
        
        # Generate default filename
        return f"document_{int(datetime.utcnow().timestamp())}"
    
    def _detect_document_type(self, content_type: str, filename: str, file_data: bytes) -> DocumentType:
        """Detect document type from various sources"""
        # Try content type first
        if content_type in self.supported_types:
            return self.supported_types[content_type]
        
        # Try file extension
        ext = Path(filename).suffix.lower()
        if ext == '.pdf':
            return DocumentType.PDF
        elif ext in ['.docx', '.doc']:
            return DocumentType.DOCX
        elif ext == '.txt':
            return DocumentType.TXT
        elif ext in ['.eml', '.email']:
            return DocumentType.EMAIL
        
        # Try magic number detection
        try:
            mime_type = magic.from_buffer(file_data, mime=True)
            if mime_type in self.supported_types:
                return self.supported_types[mime_type]
        except:
            pass
        
        # Default fallback
        logger.warning(f"Could not detect type for {filename}, defaulting to UNKNOWN")
        return DocumentType.UNKNOWN
    
    def _generate_document_id(self, blob_url: str, file_data: bytes) -> str:
        """Generate unique document ID"""
        content_hash = hashlib.sha256(file_data).hexdigest()[:16]
        url_hash = hashlib.md5(blob_url.encode()).hexdigest()[:8]
        return f"doc_{url_hash}_{content_hash}"
    
    async def _process_pdf(self, file_data: bytes) -> tuple[str, Dict[str, Any]]:
        """Process PDF document"""
        try:
            # Use PyMuPDF for better text extraction
            with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp_file:
                tmp_file.write(file_data)
                tmp_file.flush()
                
                doc = fitz.open(tmp_file.name)
                pages = []
                
                for page_num in range(len(doc)):
                    page = doc.load_page(page_num)
                    text = page.get_text()
                    if text.strip():
                        pages.append({
                            'page_number': page_num + 1,
                            'content': text.strip()
                        })
                
                doc.close()
                os.unlink(tmp_file.name)
                
                # Combine all text
                full_text = '\n\n'.join([page['content'] for page in pages])
                
                metadata = {
                    'total_pages': len(doc),
                    'pages_with_content': len(pages),
                    'processing_method': 'pymupdf'
                }
                
                return full_text, metadata
                
        except Exception as e:
            logger.warning(f"PyMuPDF failed, falling back to PyPDF2: {e}")
            return await self._process_pdf_fallback(file_data)
    
    async def _process_pdf_fallback(self, file_data: bytes) -> tuple[str, Dict[str, Any]]:
        """Fallback PDF processing with PyPDF2"""
        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp_file:
            tmp_file.write(file_data)
            tmp_file.flush()
            
            try:
                with open(tmp_file.name, 'rb') as file:
                    reader = PyPDF2.PdfReader(file)
                    pages = []
                    
                    for page_num, page in enumerate(reader.pages):
                        text = page.extract_text()
                        if text.strip():
                            pages.append({
                                'page_number': page_num + 1,
                                'content': text.strip()
                            })
                    
                    full_text = '\n\n'.join([page['content'] for page in pages])
                    
                    metadata = {
                        'total_pages': len(reader.pages),
                        'pages_with_content': len(pages),
                        'processing_method': 'pypdf2'
                    }
                    
                    return full_text, metadata
                    
            finally:
                os.unlink(tmp_file.name)
    
    async def _process_docx(self, file_data: bytes) -> tuple[str, Dict[str, Any]]:
        """Process DOCX document"""
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp_file:
            tmp_file.write(file_data)
            tmp_file.flush()
            
            try:
                doc = Document(tmp_file.name)
                paragraphs = []
                
                for para in doc.paragraphs:
                    if para.text.strip():
                        paragraphs.append(para.text.strip())
                
                full_text = '\n\n'.join(paragraphs)
                
                metadata = {
                    'total_paragraphs': len(doc.paragraphs),
                    'paragraphs_with_content': len(paragraphs),
                    'processing_method': 'python-docx'
                }
                
                return full_text, metadata
                
            finally:
                os.unlink(tmp_file.name)
    
    async def _process_email(self, file_data: bytes) -> tuple[str, Dict[str, Any]]:
        """Process email document"""
        try:
            email_content = file_data.decode('utf-8', errors='ignore')
            msg = message_from_string(email_content)
            
            # Extract headers
            subject = self._decode_header_value(msg.get('Subject', ''))
            sender = self._decode_header_value(msg.get('From', ''))
            recipient = self._decode_header_value(msg.get('To', ''))
            date = msg.get('Date', '')
            
            # Extract body
            body = self._extract_email_body(msg)
            
            # Combine content
            full_text = f"Subject: {subject}\nFrom: {sender}\nTo: {recipient}\nDate: {date}\n\n{body}"
            
            metadata = {
                'subject': subject,
                'from': sender,
                'to': recipient,
                'date': date,
                'processing_method': 'email-parser'
            }
            
            return full_text, metadata
            
        except Exception as e:
            logger.error(f"Error processing email: {e}")
            # Fallback to raw text
            return file_data.decode('utf-8', errors='ignore'), {'processing_method': 'raw'}
    
    async def _process_text(self, file_data: bytes) -> tuple[str, Dict[str, Any]]:
        """Process plain text document"""
        try:
            text = file_data.decode('utf-8')
        except UnicodeDecodeError:
            # Try other encodings
            for encoding in ['latin-1', 'cp1252', 'iso-8859-1']:
                try:
                    text = file_data.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                text = file_data.decode('utf-8', errors='ignore')
        
        metadata = {
            'character_count': len(text),
            'line_count': text.count('\n'),
            'processing_method': 'text'
        }
        
        return text, metadata
    
    def _decode_header_value(self, header_value: str) -> str:
        """Decode email header value"""
        if not header_value:
            return ''
        
        try:
            decoded_parts = decode_header(header_value)
            result = ''
            for part, encoding in decoded_parts:
                if isinstance(part, bytes):
                    if encoding:
                        result += part.decode(encoding)
                    else:
                        result += part.decode('utf-8', errors='ignore')
                else:
                    result += part
            return result
        except:
            return header_value
    
    def _extract_email_body(self, msg) -> str:
        """Extract body text from email message"""
        body = ''
        
        if msg.is_multipart():
            for part in msg.walk():
                content_type = part.get_content_type()
                if content_type == 'text/plain':
                    payload = part.get_payload(decode=True)
                    if payload:
                        body += payload.decode('utf-8', errors='ignore')
                elif content_type == 'text/html' and not body:
                    # Use HTML as fallback if no plain text
                    payload = part.get_payload(decode=True)
                    if payload:
                        from bs4 import BeautifulSoup
                        soup = BeautifulSoup(payload, 'html.parser')
                        body += soup.get_text()
        else:
            payload = msg.get_payload(decode=True)
            if payload:
                body += payload.decode('utf-8', errors='ignore')
        
        return body.strip()
    
    async def _create_chunks(self, content: str, document_id: str, metadata: Dict[str, Any]) -> List[DocumentChunk]:
        """Create text chunks from document content"""
        chunks = []
        
        # Clean and preprocess text
        cleaned_content = self.text_processor.clean_text(content)
        
        # Split into chunks
        chunk_texts = self.text_processor.split_text(
            cleaned_content,
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP
        )
        
        for i, chunk_text in enumerate(chunk_texts):
            chunk_id = f"{document_id}_chunk_{i}"
            
            chunk = DocumentChunk(
                chunk_id=chunk_id,
                content=chunk_text,
                page_number=None,  # Could be enhanced to track page numbers
                section=None,      # Could be enhanced to detect sections
                metadata={
                    'chunk_index': i,
                    'total_chunks': len(chunk_texts),
                    'character_count': len(chunk_text),
                    **metadata
                }
            )
            chunks.append(chunk)
        
        return chunks
    
    async def get_document_info(self, document_id: str) -> Optional[DocumentInfo]:
        """Get cached document info by ID"""
        # This would typically query a database
        # For now, return None as documents are processed on-demand
        return None