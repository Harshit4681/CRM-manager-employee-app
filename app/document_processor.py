import requests
import PyPDF2
import docx
import email
import io
import logging
from typing import List, Dict, Any
from urllib.parse import urlparse
from bs4 import BeautifulSoup
import aiofiles
import asyncio
from langchain.text_splitter import RecursiveCharacterTextSplitter
from app.config import settings

logger = logging.getLogger(__name__)

class DocumentProcessor:
    """
    Process various document types including PDFs, DOCX, and emails
    """
    
    def __init__(self):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP,
            length_function=len,
        )
    
    async def process_document(self, document_url: str) -> Dict[str, Any]:
        """
        Process a document from URL and extract text content
        
        Args:
            document_url: URL of the document to process
            
        Returns:
            Dict containing processed document data
        """
        try:
            logger.info(f"Processing document from URL: {document_url}")
            
            # Download document
            content = await self._download_document(document_url)
            
            # Determine document type and extract text
            doc_type = self._determine_document_type(document_url, content)
            raw_text = await self._extract_text(content, doc_type)
            
            # Split text into chunks
            chunks = self.text_splitter.split_text(raw_text)
            
            # Create structured document data
            processed_doc = {
                "url": document_url,
                "type": doc_type,
                "raw_text": raw_text,
                "chunks": chunks,
                "metadata": {
                    "chunk_count": len(chunks),
                    "total_length": len(raw_text),
                    "avg_chunk_size": len(raw_text) // len(chunks) if chunks else 0
                }
            }
            
            logger.info(f"Document processed successfully: {len(chunks)} chunks created")
            return processed_doc
            
        except Exception as e:
            logger.error(f"Error processing document: {str(e)}")
            raise Exception(f"Document processing failed: {str(e)}")
    
    async def _download_document(self, url: str) -> bytes:
        """Download document from URL"""
        try:
            response = requests.get(url, timeout=30)
            response.raise_for_status()
            return response.content
        except Exception as e:
            raise Exception(f"Failed to download document: {str(e)}")
    
    def _determine_document_type(self, url: str, content: bytes) -> str:
        """Determine document type from URL and content"""
        url_lower = url.lower()
        
        if '.pdf' in url_lower or content.startswith(b'%PDF'):
            return 'pdf'
        elif '.docx' in url_lower or b'PK' in content[:2]:
            return 'docx'
        elif '.doc' in url_lower:
            return 'doc'
        elif '@' in url.lower() or 'message-id' in str(content).lower():
            return 'email'
        else:
            # Try to detect based on content
            try:
                content.decode('utf-8')
                return 'text'
            except:
                return 'unknown'
    
    async def _extract_text(self, content: bytes, doc_type: str) -> str:
        """Extract text from document based on type"""
        try:
            if doc_type == 'pdf':
                return await self._extract_pdf_text(content)
            elif doc_type == 'docx':
                return await self._extract_docx_text(content)
            elif doc_type == 'email':
                return await self._extract_email_text(content)
            elif doc_type == 'text':
                return content.decode('utf-8')
            else:
                # Fallback: try to extract any readable text
                return content.decode('utf-8', errors='ignore')
        except Exception as e:
            logger.error(f"Error extracting text from {doc_type}: {str(e)}")
            return content.decode('utf-8', errors='ignore')
    
    async def _extract_pdf_text(self, content: bytes) -> str:
        """Extract text from PDF content"""
        try:
            pdf_file = io.BytesIO(content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
            
            return text.strip()
        except Exception as e:
            logger.error(f"Error extracting PDF text: {str(e)}")
            raise Exception(f"PDF text extraction failed: {str(e)}")
    
    async def _extract_docx_text(self, content: bytes) -> str:
        """Extract text from DOCX content"""
        try:
            doc_file = io.BytesIO(content)
            doc = docx.Document(doc_file)
            
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
            
            return text.strip()
        except Exception as e:
            logger.error(f"Error extracting DOCX text: {str(e)}")
            raise Exception(f"DOCX text extraction failed: {str(e)}")
    
    async def _extract_email_text(self, content: bytes) -> str:
        """Extract text from email content"""
        try:
            email_message = email.message_from_bytes(content)
            
            # Extract email headers
            text = f"Subject: {email_message.get('Subject', 'N/A')}\n"
            text += f"From: {email_message.get('From', 'N/A')}\n"
            text += f"To: {email_message.get('To', 'N/A')}\n"
            text += f"Date: {email_message.get('Date', 'N/A')}\n\n"
            
            # Extract email body
            if email_message.is_multipart():
                for part in email_message.walk():
                    if part.get_content_type() == "text/plain":
                        text += part.get_payload(decode=True).decode('utf-8', errors='ignore')
                    elif part.get_content_type() == "text/html":
                        html_content = part.get_payload(decode=True).decode('utf-8', errors='ignore')
                        # Parse HTML and extract text
                        soup = BeautifulSoup(html_content, 'html.parser')
                        text += soup.get_text()
            else:
                text += email_message.get_payload(decode=True).decode('utf-8', errors='ignore')
            
            return text.strip()
        except Exception as e:
            logger.error(f"Error extracting email text: {str(e)}")
            raise Exception(f"Email text extraction failed: {str(e)}")
    
    def get_document_metadata(self, processed_doc: Dict[str, Any]) -> Dict[str, Any]:
        """Extract metadata from processed document"""
        return {
            "type": processed_doc.get("type"),
            "chunk_count": len(processed_doc.get("chunks", [])),
            "total_length": len(processed_doc.get("raw_text", "")),
            "url": processed_doc.get("url")
        }